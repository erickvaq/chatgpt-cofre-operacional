# -*- coding: utf-8 -*-
"""Testes de regressao das correcoes de contaminacao e 'Sem boleto' (v2.06).

Rodar:
    WideAPP_EXTRA/.venv/Scripts/python.exe -m pytest WideAPP_EXTRA/tests/test_isolamento_regressao.py -q
ou simplesmente:
    WideAPP_EXTRA/.venv/Scripts/python.exe WideAPP_EXTRA/tests/test_isolamento_regressao.py
"""
import json
import sys
import tempfile
from pathlib import Path

APP_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_ROOT))

from app.isolamento_cliente import (
    mesma_identidade,
    isolar_registros,
    deduplicar_cobrancas,
)


# Reproduz a contaminacao real: o arquivo de ANA CLEIDE continha 72 cobrancas
# de tres clientes "...SANTANA" porque "ANA" e substring de "SANTANA".
FIXTURE_CONTAMINADO = {
    "cliente": "ANA CLEIDE DOS SANTOS DIAS",
    "carnes": [
        {"carne": "216", "cliente": "ANA CLEIDE DOS SANTOS DIAS", "referencia": "lt e5",
         "valor_parcela": 93.0, "status": "Pendente"},
    ],
    "cobrancas": [
        {"id": "4774", "cliente": "ANA CLEIDE DOS SANTOS DIAS", "descricao": "lt e5",
         "valor_original": 93.0, "valor_recebido": 93.0, "status": "Recebido"},
        {"id": "4773", "cliente": "ANA CLEIDE DOS SANTOS DIAS", "descricao": "lt e5",
         "valor_original": 93.0, "valor_recebido": 93.0, "status": "Recebido"},
        {"id": "9001", "cliente": "Carlos Alberto Santana", "descricao": "carne x",
         "valor_original": 100.0, "valor_recebido": 100.0, "status": "Recebido"},
        {"id": "9002", "cliente": "Eliel Hora Santana", "descricao": "carne y",
         "valor_original": 100.0, "valor_recebido": 100.0, "status": "Recebido"},
        {"id": "9003", "cliente": "Marinalva Hora Santana", "descricao": "carne z",
         "valor_original": 100.0, "valor_recebido": 100.0, "status": "Recebido"},
    ],
}


def test_identidade_nao_casa_ana_com_santana():
    alvo = "ANA CLEIDE DOS SANTOS DIAS"
    assert mesma_identidade(alvo, "ANA CLEIDE DOS SANTOS DIAS") is True
    assert mesma_identidade(alvo, "Ana Cleide") is True  # WidePay abreviado
    assert mesma_identidade(alvo, "Carlos Alberto Santana") is False
    assert mesma_identidade(alvo, "Eliel Hora Santana") is False
    assert mesma_identidade(alvo, "Marinalva Hora Santana") is False


def test_isolamento_remove_terceiros():
    iso = isolar_registros(FIXTURE_CONTAMINADO["cobrancas"], "ANA CLEIDE DOS SANTOS DIAS")
    donos = {r.get("cliente") for r in iso["mantidos"]}
    assert donos == {"ANA CLEIDE DOS SANTOS DIAS"}, f"vazou terceiro: {donos}"
    assert len(iso["mantidos"]) == 2
    assert len(iso["descartados"]) == 3
    proprietarios_descartados = {d["proprietario"] for d in iso["descartados"]}
    assert proprietarios_descartados == {
        "Carlos Alberto Santana", "Eliel Hora Santana", "Marinalva Hora Santana"
    }


def test_dedup_por_id():
    cobr = [
        {"id": "10", "descricao": "a", "valor_original": 1, "valor_recebido": 1, "status": "Recebido"},
        {"id": "10", "descricao": "a", "valor_original": 1, "valor_recebido": 1, "status": "Recebido"},
        {"id": "11", "descricao": "b", "valor_original": 1, "valor_recebido": 1, "status": "Recebido"},
    ]
    unicos, removidos = deduplicar_cobrancas(cobr)
    assert len(unicos) == 2
    assert removidos == 1


def test_gerar_relatorio_excel_isola_e_contrato_nao_confirmado():
    """Integra o gerador de XLSX: com contrato 0 parcelas, nao deve inflar
    parcelas com terceiros nem afirmar quitacao."""
    from app.gerar_relatorio_excel import gerar_relatorio_excel
    dados_cliente = {"cliente": "ANA CLEIDE DOS SANTOS DIAS", "lote": "E5", "quadra": "E"}
    with tempfile.TemporaryDirectory() as td:
        saida = Path(td) / "rel.xlsx"
        gerar_relatorio_excel(
            dict(FIXTURE_CONTAMINADO), dados_cliente,
            valor_parcela=93.0, total_parcelas_contrato=0,
            valor_total_contrato=0.0, caminho_saida=str(saida),
        )
        assert saida.exists()
        from openpyxl import load_workbook
        wb = load_workbook(saida)
        ws = wb["Resumo"]
        textos = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        blob = " | ".join(textos)
        # Nenhum nome de terceiro pode aparecer no relatorio de Ana Cleide
        for terceiro in ("Carlos Alberto Santana", "Eliel Hora Santana", "Marinalva Hora Santana"):
            assert terceiro not in blob, f"terceiro vazou no XLSX: {terceiro}"
        # Contrato nao confirmado deve ser sinalizado
        assert any("CONTRATO NAO CONFIRMADO" in t for t in textos), "faltou rotulo de contrato nao confirmado"
        # Aba Pagamentos deve conter apenas 2 recebidos de Ana Cleide (nao 5)
        ws_pag = wb["Pagamentos Recebidos"]
        linhas_dados = ws_pag.max_row - 2  # header + linha total
        assert linhas_dados == 2, f"esperado 2 pagamentos de Ana Cleide, veio {linhas_dados}"


def test_cache_sem_boleto_lote_vazio():
    """Rafaela de Jesus Cruz: cobrancas com lote em branco no WidePay nao
    podem faze-la aparecer como 'Sem boleto' quando o indice informa lote."""
    from app import widepay_boletos_cache as wc
    if not wc.config.WIDEPAY_BOLETOS_CACHE_JSON.exists():
        print("SKIP: cache real ausente")
        return
    raw = wc.montar_raw_cliente("RAFAELA DE JESUS CRUZ", "12", "G")
    assert raw is not None, "Rafaela voltou a aparecer como 'Sem boleto'"
    donos = {c.get("cliente") for c in raw["cobrancas"]}
    # nenhum terceiro
    for d in donos:
        assert mesma_identidade("RAFAELA DE JESUS CRUZ", d), f"terceiro no cache de Rafaela: {d}"


def _run_all():
    falhas = 0
    for nome, fn in sorted(globals().items()):
        if nome.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"[PASS] {nome}")
            except AssertionError as e:
                falhas += 1
                print(f"[FAIL] {nome}: {e}")
            except Exception as e:
                falhas += 1
                print(f"[ERRO] {nome}: {type(e).__name__}: {e}")
    print(f"\n{'TODOS OK' if falhas == 0 else str(falhas) + ' FALHA(S)'}")
    return falhas


if __name__ == "__main__":
    sys.exit(1 if _run_all() else 0)
